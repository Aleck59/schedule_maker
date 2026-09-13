"""Загрузка справочников из таблицы.

Один файл — один тип данных, колонки подписаны по-русски. Так заполнять
первичные данные быстрее, чем вбивать руками в формы.
"""

from __future__ import annotations

import io
from typing import Any, ClassVar

from sqlalchemy import select
from sqlalchemy.orm import Session

from schedule_maker.enums import DeliveryMode, RoomKind, StudyForm
from schedule_maker.models import (
    Campus,
    Faculty,
    Room,
    RoomFeature,
    StudentGroup,
    Subgroup,
    Subject,
    Teacher,
)
from schedule_maker.plugins.api import ImporterPlugin, ImportResult, PluginManifest

ROOM_KIND_BY_NAME = {
    "лекционная": RoomKind.LECTURE_HALL,
    "поточная": RoomKind.LECTURE_HALL,
    "семинарская": RoomKind.SEMINAR,
    "лаборатория": RoomKind.LAB,
    "компьютерный класс": RoomKind.COMPUTER,
    "спортзал": RoomKind.GYM,
}
STUDY_FORM_BY_NAME = {
    "очная": StudyForm.FULL_TIME,
    "очно-заочная": StudyForm.EVENING,
    "вечерняя": StudyForm.EVENING,
    "заочная": StudyForm.EXTRAMURAL,
}


def _rows(raw: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(raw), data_only=True)
    sheet = book.active
    if sheet is None:  # pragma: no cover - книга без листов
        return []
    return [
        list(row)
        for row in sheet.iter_rows(values_only=True)
        if any(value is not None for value in row)
    ]


def _features(session: Session, text: str) -> list[RoomFeature]:
    """Разобрать колонку «Оборудование» и подтянуть строки справочника.

    Названия, которых в справочнике ещё нет, заводятся на месте: иначе
    импорт молча терял бы половину колонки. Регистр не различается —
    «Проектор» и «проектор» должны быть одной строкой, а не двумя.
    """
    names = [part.strip() for part in text.replace(";", ",").split(",") if part.strip()]
    if not names:
        return []
    known = {feature.name.casefold(): feature for feature in session.scalars(select(RoomFeature))}
    chosen: dict[int | str, RoomFeature] = {}
    for name in names:
        feature = known.get(name.casefold())
        if feature is None:
            feature = RoomFeature(name=name)
            session.add(feature)
            session.flush()
            known[name.casefold()] = feature
        chosen[feature.id] = feature
    return list(chosen.values())


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _slug(value: str, prefix: str) -> str:
    from schedule_maker.web.routers.admin_catalog import _slugify

    return _slugify(value) or prefix


class TeacherImporter(ImporterPlugin):
    """Колонки: ФИО, Кафедра, Почта, Формат, Пар в день, Пар в неделю."""

    key: ClassVar[str] = "import.teachers_xlsx"
    title: ClassVar[str] = "Преподаватели из XLSX"
    accepts: ClassVar[str] = ".xlsx"
    manifest: ClassVar[PluginManifest | None] = PluginManifest(
        key="import.teachers_xlsx",
        name="Загрузка преподавателей",
        description="Колонки: ФИО, Кафедра, Почта, Формат, Пар в день, Пар в неделю.",
        kind="importer",
        builtin=True,
    )

    def run(self, session: Session, raw: bytes, options: dict[str, Any]) -> ImportResult:
        result = ImportResult()
        rows = _rows(raw)
        for number, row in enumerate(rows[1:], start=2):
            full_name = _text(row[0] if row else "")
            if not full_name:
                result.skipped += 1
                continue
            teacher = session.scalars(select(Teacher).where(Teacher.full_name == full_name)).first()
            created = teacher is None
            if teacher is None:
                teacher = Teacher(full_name=full_name, slug=_slug(full_name, f"teacher-{number}"))
                session.add(teacher)
            teacher.department = _text(row[1]) if len(row) > 1 else teacher.department
            teacher.email = _text(row[2]) if len(row) > 2 else teacher.email
            mode = _text(row[3]).lower() if len(row) > 3 else ""
            teacher.delivery_mode = (
                DeliveryMode.ONLINE if mode.startswith("онлайн") else DeliveryMode.OFFLINE
            )
            if len(row) > 4 and row[4]:
                teacher.max_pairs_per_day = int(row[4])
            if len(row) > 5 and row[5]:
                teacher.max_pairs_per_week = int(row[5])
            session.flush()
            if created:
                result.created += 1
            else:
                result.updated += 1
        return result


class GroupImporter(ImporterPlugin):
    """Колонки: Название, Курс, Направление, Филиал, Форма, Численность, Подгрупп."""

    key: ClassVar[str] = "import.groups_xlsx"
    title: ClassVar[str] = "Группы из XLSX"
    accepts: ClassVar[str] = ".xlsx"
    manifest: ClassVar[PluginManifest | None] = PluginManifest(
        key="import.groups_xlsx",
        name="Загрузка групп",
        description=(
            "Колонки: Название, Курс, Направление, Филиал, Форма обучения, "
            "Численность, Число подгрупп (0 — без деления)."
        ),
        kind="importer",
        builtin=True,
    )

    def run(self, session: Session, raw: bytes, options: dict[str, Any]) -> ImportResult:
        result = ImportResult()
        for number, row in enumerate(_rows(raw)[1:], start=2):
            name = _text(row[0] if row else "")
            if not name:
                result.skipped += 1
                continue
            faculty_name = _text(row[2]) if len(row) > 2 else ""
            campus_name = _text(row[3]) if len(row) > 3 else ""
            faculty = session.scalars(select(Faculty).where(Faculty.name == faculty_name)).first()
            campus = session.scalars(select(Campus).where(Campus.name == campus_name)).first()
            if faculty is None or campus is None:
                result.errors.append(
                    f"Строка {number}: не найдено направление «{faculty_name}» "
                    f"или филиал «{campus_name}»"
                )
                continue

            group = session.scalars(select(StudentGroup).where(StudentGroup.name == name)).first()
            created = group is None
            if group is None:
                group = StudentGroup(name=name, slug=_slug(name, f"group-{number}"))
                session.add(group)
            group.course = int(row[1]) if len(row) > 1 and row[1] else 1
            group.faculty_id = faculty.id
            group.campus_id = campus.id
            form = _text(row[4]).lower() if len(row) > 4 else ""
            group.study_form = STUDY_FORM_BY_NAME.get(form, StudyForm.FULL_TIME)
            group.size = int(row[5]) if len(row) > 5 and row[5] else 25
            subgroups = int(row[6]) if len(row) > 6 and row[6] else 0
            group.split_flag = subgroups > 1
            group.subgroup_count = max(1, subgroups)
            session.flush()

            existing = {s.index for s in group.subgroups}
            if group.split_flag:
                for index in range(1, group.subgroup_count + 1):
                    if index not in existing:
                        session.add(
                            Subgroup(
                                group_id=group.id,
                                index=index,
                                size=max(1, group.size // group.subgroup_count),
                            )
                        )
            else:
                for subgroup in list(group.subgroups):
                    session.delete(subgroup)
            session.flush()
            if created:
                result.created += 1
            else:
                result.updated += 1
        return result


class RoomImporter(ImporterPlugin):
    """Колонки: Номер, Название, Филиал, Тип, Мест, Оборудование."""

    key: ClassVar[str] = "import.rooms_xlsx"
    title: ClassVar[str] = "Аудитории из XLSX"
    accepts: ClassVar[str] = ".xlsx"
    manifest: ClassVar[PluginManifest | None] = PluginManifest(
        key="import.rooms_xlsx",
        name="Загрузка аудиторий",
        description="Колонки: Номер, Название, Филиал, Тип, Мест, Оборудование.",
        kind="importer",
        builtin=True,
    )

    def run(self, session: Session, raw: bytes, options: dict[str, Any]) -> ImportResult:
        result = ImportResult()
        for number, row in enumerate(_rows(raw)[1:], start=2):
            code = _text(row[0] if row else "")
            campus_name = _text(row[2]) if len(row) > 2 else ""
            campus = session.scalars(select(Campus).where(Campus.name == campus_name)).first()
            if not code or campus is None:
                result.errors.append(f"Строка {number}: нет номера или филиала «{campus_name}»")
                continue
            room = session.scalars(
                select(Room).where(Room.code == code, Room.campus_id == campus.id)
            ).first()
            created = room is None
            if room is None:
                room = Room(code=code, campus_id=campus.id)
                session.add(room)
            room.name = _text(row[1]) if len(row) > 1 else ""
            room.kind = ROOM_KIND_BY_NAME.get(
                _text(row[3]).lower() if len(row) > 3 else "", RoomKind.SEMINAR
            )
            room.capacity = int(row[4]) if len(row) > 4 and row[4] else 30
            room.features = _features(session, _text(row[5]) if len(row) > 5 else "")
            session.flush()
            if created:
                result.created += 1
            else:
                result.updated += 1
        return result


class SubjectImporter(ImporterPlugin):
    """Колонки: Название, Сокращение."""

    key: ClassVar[str] = "import.subjects_xlsx"
    title: ClassVar[str] = "Дисциплины из XLSX"
    accepts: ClassVar[str] = ".xlsx"
    manifest: ClassVar[PluginManifest | None] = PluginManifest(
        key="import.subjects_xlsx",
        name="Загрузка дисциплин",
        description="Колонки: Название, Сокращение.",
        kind="importer",
        builtin=True,
    )

    def run(self, session: Session, raw: bytes, options: dict[str, Any]) -> ImportResult:
        result = ImportResult()
        for row in _rows(raw)[1:]:
            name = _text(row[0] if row else "")
            if not name:
                result.skipped += 1
                continue
            subject = session.scalars(select(Subject).where(Subject.name == name)).first()
            if subject is None:
                session.add(Subject(name=name, short=_text(row[1]) if len(row) > 1 else ""))
                result.created += 1
            else:
                if len(row) > 1 and _text(row[1]):
                    subject.short = _text(row[1])
                result.updated += 1
        session.flush()
        return result


PLUGINS = [TeacherImporter, GroupImporter, RoomImporter, SubjectImporter]
