"""Импорт нагрузки из Excel или CSV.

Учебный план и нагрузка почти всегда уже существуют в виде таблицы, и
перенабирать их руками — самая обидная часть работы. Импортёр читает файл как
есть: строка = одно требование («столько-то пар такой-то дисциплины такой-то
группе с таким-то преподавателем»).

Соответствие колонок угадывается по заголовкам, но остаётся настраиваемым:
в каждом вузе таблица своя, и подгонять её под программу неправильно —
программа должна подстроиться сама.
"""

from __future__ import annotations

import csv
import io
from typing import Any, ClassVar

from ...domain.models import (
    Discipline,
    Lesson,
    LessonType,
    Location,
    Problem,
    RoomKind,
    StudentGroup,
    Teacher,
)
from ...domain.timegrid import default_period_templates
from ...plugins.api import ImportBatch, ImportPreview

#: Поле → слова, по которым его заголовок узнаётся в чужой таблице.
FIELD_HINTS: dict[str, tuple[str, ...]] = {
    "group": ("группа", "групп"),
    "course": ("курс",),
    "program": ("направление", "специальность", "факультет", "профиль"),
    "headcount": ("человек", "числен", "студент", "кол-во студ"),
    "discipline": ("дисциплина", "предмет"),
    "department": ("кафедра",),
    "teacher": ("преподаватель", "фио", "педагог"),
    "lesson_type": ("вид занятия", "вид", "тип занятия", "тип"),
    "pairs": ("пар", "часов", "количество", "объём", "объем"),
    "room_kind": ("аудитория", "тип аудитории", "помещение"),
    "required_start": ("начало", "время"),
}

FIELD_TITLES_RU: dict[str, str] = {
    "group": "Группа",
    "course": "Курс",
    "program": "Направление",
    "headcount": "Человек в группе",
    "discipline": "Дисциплина",
    "department": "Кафедра",
    "teacher": "Преподаватель",
    "lesson_type": "Вид занятия",
    "pairs": "Количество пар",
    "room_kind": "Тип аудитории",
    "required_start": "Время начала",
}

REQUIRED_FIELDS = ("group", "discipline", "teacher", "pairs")

LESSON_TYPE_WORDS: dict[str, LessonType] = {
    "лекц": LessonType.LECTURE,
    "практ": LessonType.PRACTICE,
    "семинар": LessonType.PRACTICE,
    "лаб": LessonType.LAB,
}

ROOM_KIND_WORDS: dict[str, RoomKind] = {
    "лекц": RoomKind.LECTURE,
    "лаб": RoomKind.LAB,
    "комп": RoomKind.COMPUTER,
    "ауд": RoomKind.PRACTICE,
}


class ExcelImporter:
    id: ClassVar[str] = "excel"
    title: ClassVar[str] = "Нагрузка из Excel / CSV"
    accepts: ClassVar[tuple[str, ...]] = (".xlsx", ".xlsm", ".csv")

    # --- разбор файла --------------------------------------------------

    def preview(self, data: bytes, mapping: dict[str, str] | None = None) -> ImportPreview:
        """Показать, что программа поняла, — до того как что-то менять."""
        columns, rows = _read_table(data)
        detected = mapping or detect_mapping(columns)
        problems = _validate_mapping(detected, columns)
        return ImportPreview(
            columns=columns,
            rows=[{k: str(v) for k, v in row.items()} for row in rows[:15]],
            problems=problems,
            detected_mapping=detected,
        )

    def load(self, data: bytes, mapping: dict[str, str] | None = None) -> ImportBatch:
        columns, rows = _read_table(data)
        mapping = mapping or detect_mapping(columns)
        problems = _validate_mapping(mapping, columns)
        if problems:
            return ImportBatch(problem=Problem(), problems=problems)

        location = Location(id=1, city="Филиал", name="")
        groups: dict[str, StudentGroup] = {}
        teachers: dict[str, Teacher] = {}
        disciplines: dict[str, Discipline] = {}
        lessons: list[Lesson] = []

        for number, row in enumerate(rows, start=2):
            group_name = _text(row, mapping, "group")
            discipline_name = _text(row, mapping, "discipline")
            teacher_name = _text(row, mapping, "teacher")
            pairs = _int(row, mapping, "pairs")

            if not (group_name and discipline_name and teacher_name):
                continue
            if pairs <= 0:
                problems.append(f"Строка {number}: не указано количество пар — строка пропущена.")
                continue

            group = groups.get(group_name)
            if group is None:
                group = StudentGroup(
                    id=len(groups) + 1,
                    name=group_name,
                    course=_int(row, mapping, "course") or 1,
                    program=_text(row, mapping, "program"),
                    headcount=_int(row, mapping, "headcount") or 25,
                    location_id=1,
                )
                groups[group_name] = group

            teacher = teachers.get(teacher_name)
            if teacher is None:
                teacher = Teacher(
                    id=len(teachers) + 1,
                    full_name=teacher_name,
                    department=_text(row, mapping, "department"),
                )
                teachers[teacher_name] = teacher

            discipline = disciplines.get(discipline_name)
            if discipline is None:
                discipline = Discipline(
                    id=len(disciplines) + 1,
                    name=discipline_name,
                    department=_text(row, mapping, "department"),
                )
                disciplines[discipline_name] = discipline

            lessons.append(
                Lesson(
                    id=len(lessons) + 1,
                    discipline_id=discipline.id,
                    group_id=group.id,
                    teacher_id=teacher.id,
                    location_id=1,
                    lesson_type=_lesson_type(_text(row, mapping, "lesson_type")),
                    pairs_total=pairs,
                    required_room_kind=_room_kind(_text(row, mapping, "room_kind")),
                    required_start=_time(_text(row, mapping, "required_start")),
                )
            )

        problem = Problem(
            locations=[location],
            groups=list(groups.values()),
            teachers=list(teachers.values()),
            disciplines=list(disciplines.values()),
            lessons=lessons,
            period_templates=default_period_templates(1),
        )
        return ImportBatch(
            problem=problem,
            created={
                "группы": len(groups),
                "преподаватели": len(teachers),
                "дисциплины": len(disciplines),
                "занятия": len(lessons),
            },
            problems=problems,
        )


# ----------------------------------------------------------------------


def _read_table(data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Прочитать .xlsx или .csv в одинаковый вид: заголовки и строки-словари."""
    if data[:2] == b"PK":  # xlsx — это zip-архив
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = next(it, None) or ()
        columns = [str(c).strip() if c is not None else "" for c in header]
        rows = []
        for values in it:
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append({columns[i]: values[i] for i in range(min(len(columns), len(values)))})
        return columns, rows

    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    columns = [c.strip() for c in (reader.fieldnames or [])]
    rows = [dict(r) for r in reader]
    return columns, rows


def detect_mapping(columns: list[str]) -> dict[str, str]:
    """Угадать, какая колонка чему соответствует, по словам в заголовке."""
    mapping: dict[str, str] = {}
    taken: set[str] = set()
    for field, hints in FIELD_HINTS.items():
        for column in columns:
            if column in taken or not column:
                continue
            low = column.strip().lower()
            if any(h in low for h in hints):
                mapping[field] = column
                taken.add(column)
                break
    return mapping


def _validate_mapping(mapping: dict[str, str], columns: list[str]) -> list[str]:
    problems = []
    for field in REQUIRED_FIELDS:
        if not mapping.get(field):
            problems.append(
                f"Не найдена колонка «{FIELD_TITLES_RU[field]}». "
                f"Укажите её вручную. Колонки файла: {', '.join(c for c in columns if c)}."
            )
    return problems


def _text(row: dict, mapping: dict[str, str], field: str) -> str:
    column = mapping.get(field)
    if not column:
        return ""
    value = row.get(column)
    return "" if value is None else str(value).strip()


def _int(row: dict, mapping: dict[str, str], field: str) -> int:
    raw = _text(row, mapping, field).replace(",", ".")
    if not raw:
        return 0
    try:
        return int(float(raw))
    except ValueError:
        return 0


def _lesson_type(raw: str) -> LessonType:
    low = raw.lower()
    for word, value in LESSON_TYPE_WORDS.items():
        if word in low:
            return value
    return LessonType.PRACTICE


def _room_kind(raw: str) -> RoomKind | None:
    low = raw.lower()
    for word, value in ROOM_KIND_WORDS.items():
        if word in low:
            return value
    return None


def _time(raw: str):
    """«16:00» или «16.00» — в объект времени; всё непонятное игнорируется."""
    from datetime import time

    cleaned = raw.replace(".", ":").strip()
    if not cleaned or ":" not in cleaned:
        return None
    try:
        hours, minutes = cleaned.split(":")[:2]
        return time(int(hours), int(minutes))
    except ValueError:
        return None


def register() -> list[ExcelImporter]:
    return [ExcelImporter()]
