"""Выгрузка расписания в Excel.

Учебная часть всё равно живёт в Excel: пересылает, печатает, правит. Лист
повторяет привычную сетку — строки это пары, столбцы дни недели.
"""

from __future__ import annotations

from typing import ClassVar

from ...domain.timegrid import DAY_FULL_RU
from ...domain.timetable import Timetable
from ...plugins.api import ExportView
from ...web.view import build_grid


class XlsxExporter:
    id: ClassVar[str] = "xlsx"
    title: ClassVar[str] = "Excel (.xlsx)"
    media_type: ClassVar[str] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension: ClassVar[str] = "xlsx"

    def export(self, tt: Timetable, view: ExportView) -> bytes:
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        grid = build_grid(tt, kind=view.kind, subject_id=view.subject_id, title=view.title)

        wb = Workbook()
        ws = wb.active
        ws.title = "Расписание"

        thin = Side(style="thin", color="D0D4DA")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="EEF2F8")
        wrap = Alignment(wrap_text=True, vertical="top")

        ws.cell(row=1, column=1, value=view.title).font = Font(bold=True, size=14)
        ws.cell(row=3, column=1, value="Пара").font = Font(bold=True)
        for i, day in enumerate(DAY_FULL_RU[: tt.problem.days]):
            cell = ws.cell(row=3, column=2 + i, value=day.capitalize())
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        for r, row in enumerate(grid.periods, start=4):
            label = f"{row.period}\n{row.time_label}" if row.time_label else str(row.period)
            head = ws.cell(row=r, column=1, value=label)
            head.alignment = wrap
            head.border = border
            head.fill = header_fill
            for c, day in enumerate(grid.day_indexes):
                pairs = grid.at(day, row.period)
                text = "\n\n".join(
                    "\n".join(
                        filter(
                            None,
                            [
                                f"{pair.discipline} ({pair.lesson_type})",
                                f"{pair.group} · {pair.teacher}",
                                "дистанционно" if pair.is_online else f"ауд. {pair.room}",
                                pair.parity,
                            ],
                        )
                    )
                    for pair in pairs
                )
                cell = ws.cell(row=r, column=2 + c, value=text)
                cell.alignment = wrap
                cell.border = border

        ws.column_dimensions["A"].width = 12
        for i in range(tt.problem.days):
            ws.column_dimensions[chr(ord("B") + i)].width = 30
        for r in range(4, 4 + len(grid.periods)):
            ws.row_dimensions[r].height = 72

        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


def register() -> list[XlsxExporter]:
    return [XlsxExporter()]
