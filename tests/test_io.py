"""Выгрузка, загрузка и внешние расписания."""

from __future__ import annotations

import io

from sqlalchemy import select
from sqlalchemy.orm import Session

from schedule_maker.models import ExternalBusy, ExternalSource, Room, StudentGroup, Teacher
from schedule_maker.plugins.builtin.io_formats.exporters import (
    CsvExporter,
    IcsExporter,
    XlsxExporter,
)
from schedule_maker.plugins.builtin.io_formats.importers import (
    GroupImporter,
    RoomImporter,
    TeacherImporter,
)
from schedule_maker.plugins.builtin.source_external import parse_ics
from schedule_maker.services.external import sync_source
from schedule_maker.services.problem_builder import build_problem, load_timetable
from schedule_maker.services.versions import working_version


def xlsx_bytes(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def filled(session: Session):
    version = working_version(session)
    problem = build_problem(session)
    return problem, load_timetable(session, version.id)


def test_выгрузка_в_календарь(demo, session: Session):
    problem, timetable = filled(session)
    from schedule_maker.domain import Placement

    timetable.add(Placement(demand_id=next(iter(problem.demands)), component=0, day=0, index=0))
    artifact = IcsExporter().export(problem, timetable, {"calendar_name": "Проверка"})
    text = artifact.data.decode("utf-8")
    assert text.startswith("BEGIN:VCALENDAR")
    assert "X-WR-CALNAME:Проверка" in text
    assert text.count("BEGIN:VEVENT") == 18  # одна пара на 18 учебных недель


def test_нечётная_неделя_даёт_вдвое_меньше_событий(demo, session: Session):
    from schedule_maker.domain import Placement
    from schedule_maker.enums import WeekParity

    problem, timetable = filled(session)
    timetable.placements.clear()
    timetable.reindex()
    timetable.add(
        Placement(
            demand_id=next(iter(problem.demands)),
            component=0,
            day=0,
            index=0,
            parity=WeekParity.ODD,
        )
    )
    text = IcsExporter().export(problem, timetable, {}).data.decode("utf-8")
    assert text.count("BEGIN:VEVENT") == 9


def test_выгрузка_в_csv(demo, session: Session):
    from schedule_maker.domain import Placement

    problem, timetable = filled(session)
    timetable.add(Placement(demand_id=next(iter(problem.demands)), component=0, day=0, index=0))
    text = CsvExporter().export(problem, timetable, {}).data.decode("utf-8")
    assert "День;Пара;Начало" in text
    assert "Понедельник" in text


def test_выгрузка_в_xlsx(demo, session: Session):
    from openpyxl import load_workbook

    problem, timetable = filled(session)
    artifact = XlsxExporter().export(problem, timetable, {})
    book = load_workbook(io.BytesIO(artifact.data))
    assert "БИО-101" in book.sheetnames
    sheet = book["БИО-101"]
    assert sheet.cell(row=1, column=2).value == "Понедельник"


def test_загрузка_преподавателей(demo, session: Session):
    raw = xlsx_bytes(
        [
            ["ФИО", "Кафедра", "Почта", "Формат", "Пар в день", "Пар в неделю"],
            ["Новиков Пётр Сергеевич", "Кафедра права", "novikov@example.edu", "офлайн", 3, 18],
            ["Магомедов Али Гаджиевич", "Новая кафедра", "", "онлайн", 5, 20],
        ]
    )
    result = TeacherImporter().run(session, raw, {})
    assert result.created == 1 and result.updated == 1
    session.flush()
    новый = session.scalars(
        select(Teacher).where(Teacher.full_name == "Новиков Пётр Сергеевич")
    ).first()
    assert новый is not None and новый.max_pairs_per_day == 3
    обновлённый = session.scalars(select(Teacher).where(Teacher.slug == "magomedov")).first()
    assert обновлённый.delivery_mode == "online"


def test_загрузка_групп_создаёт_подгруппы(demo, session: Session):
    raw = xlsx_bytes(
        [
            ["Название", "Курс", "Направление", "Филиал", "Форма", "Численность", "Подгрупп"],
            ["ЮР-301", 3, "Юриспруденция", "Махачкала", "очная", 24, 3],
            ["БИО-301", 3, "Биология", "Махачкала", "очная", 20, 0],
        ]
    )
    result = GroupImporter().run(session, raw, {})
    assert result.created == 2 and not result.errors
    session.flush()
    law = session.scalars(select(StudentGroup).where(StudentGroup.name == "ЮР-301")).first()
    bio = session.scalars(select(StudentGroup).where(StudentGroup.name == "БИО-301")).first()
    assert law.split_flag and len(law.subgroups) == 3
    assert not bio.split_flag and bio.subgroups == []


def test_загрузка_групп_сообщает_об_ошибке(demo, session: Session):
    raw = xlsx_bytes(
        [
            ["Название", "Курс", "Направление", "Филиал", "Форма", "Численность", "Подгрупп"],
            ["ХИМ-101", 1, "Химия", "Дербент", "очная", 20, 0],
        ]
    )
    result = GroupImporter().run(session, raw, {})
    assert not result.ok
    assert "Дербент" in result.errors[0]


def test_загрузка_аудиторий(demo, session: Session):
    raw = xlsx_bytes(
        [
            ["Номер", "Название", "Филиал", "Тип", "Мест", "Оборудование"],
            ["402", "Новая лаборатория", "Махачкала", "лаборатория", 18, "вытяжка"],
        ]
    )
    result = RoomImporter().run(session, raw, {})
    assert result.created == 1
    session.flush()
    room = session.scalars(select(Room).where(Room.code == "402")).first()
    assert room.kind == "lab" and room.capacity == 18


ICS = """BEGIN:VCALENDAR
BEGIN:VEVENT
DTSTART:20260914T094000
SUMMARY:Математика в колледже
ATTENDEE:MAILTO:petrova@example.edu
END:VEVENT
BEGIN:VEVENT
DTSTART:20260916T080000
SUMMARY:Информатика
ATTENDEE:MAILTO:petrova@example.edu
END:VEVENT
END:VCALENDAR
"""


def test_разбор_календаря():
    slots = parse_ics(ICS)
    assert len(slots) == 2
    assert slots[0].teacher_ref == "petrova@example.edu"
    assert slots[0].day_of_week == 0  # 14 сентября 2026 — понедельник
    assert slots[0].start_minutes == 9 * 60 + 40


def test_синхронизация_внешнего_расписания(demo, session: Session, monkeypatch):
    """Занятость в Колледже превращается в запрет на нашей сетке."""
    source = session.scalars(select(ExternalSource)).first()
    for row in list(source.busy_slots):
        session.delete(row)
    session.flush()

    from schedule_maker.plugins.registry import get_registry

    plugin = get_registry().instance("source.ics_url")
    monkeypatch.setattr(plugin, "fetch", lambda config: parse_ics(ICS))

    result = sync_source(session, source)
    assert result.imported == 2
    assert not result.unmatched
    session.flush()

    teacher = session.scalars(select(Teacher).where(Teacher.slug == "petrova")).first()
    busy = list(session.scalars(select(ExternalBusy).where(ExternalBusy.teacher_id == teacher.id)))
    assert {(b.day_of_week, b.slot_index) for b in busy} == {(0, 1), (2, 0)}


def test_несопоставленные_преподаватели_попадают_в_отчёт(demo, session: Session, monkeypatch):
    from schedule_maker.plugins.api import BusySlot
    from schedule_maker.plugins.registry import get_registry

    source = session.scalars(select(ExternalSource)).first()
    plugin = get_registry().instance("source.ics_url")
    monkeypatch.setattr(
        plugin,
        "fetch",
        lambda config: [
            BusySlot(teacher_ref="unknown@example.edu", day_of_week=0, start_minutes=480)
        ],
    )
    result = sync_source(session, source)
    assert result.imported == 0
    assert "unknown@example.edu" in result.unmatched
